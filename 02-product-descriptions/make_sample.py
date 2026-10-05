#!/usr/bin/env python3
"""Build the 20-product SAMPLE-DATA sheet for the product-descriptions template.

All brand names are INVENTED; features are plain, factual spec text so the AI never has to
invent a claim. Columns (Arabic, what an Arab seller's own sheet looks like):
    الاسم | الفئة | المميزات | السعر
The workflow adds two columns:  الوصف (عربي)  and  Description (EN).

    python3 make_sample.py <out.xlsx>
"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# name, category, features, price  (SAMPLE DATA — invented brands)
PRODUCTS = [
    ("عطر ليل العود من دار نسمات", "عطور", "عود كمبودي، مسك أبيض، يدوم حتى ٨ ساعات، حجم ٥٠ مل", "189 ريال"),
    ("كريم ترطيب نواره بزبدة الشيا", "عناية بالبشرة", "زبدة شيا وفيتامين E، خالٍ من العطور، للبشرة الجافة، ٥٠ مل", "75 ريال"),
    ("عباية كلوش رونق كريب سادة", "أزياء نسائية", "قماش كريب، أسود، مقاسات S إلى XL، قابلة للغسل في الغسالة", "240 ريال"),
    ("ساعة يد زمن بسوار جلد بني", "إكسسوارات", "ميناء كلاسيكي، سوار جلد طبيعي، مقاومة للماء ٣ أمتار، ضمان سنة", "310 ريال"),
    ("سماعة بلوتوث صدى لاسلكية", "إلكترونيات", "بلوتوث ٥.٣، عزل ضجيج، بطارية ٣٠ ساعة، شحن سريع USB-C", "155 ريال"),
    ("طقم أكواب قهوة ضيافة ٦ قطع", "المنزل والمطبخ", "سيراميك مزجج، ٦ أكواب سعة ٢٠٠ مل، آمنة في غسالة الصحون", "95 ريال"),
    ("حقيبة ظهر رحّالة مقاومة للماء", "حقائب", "بوليستر مقاوم للماء، جيب للابتوب ١٥.٦ بوصة، منفذ شحن USB", "199 ريال"),
    ("سجادة صلاة سكينة مبطّنة", "مستلزمات دينية", "إسفنج ذاكرة، مقاس ١١٠×٦٥ سم، حقيبة حمل، قابلة للغسل", "120 ريال"),
    ("لعبة تركيب خشبية نمو للأطفال", "ألعاب أطفال", "خشب زان طبيعي، ٤٨ قطعة، حواف ناعمة آمنة، من عمر ٣ سنوات", "85 ريال"),
    ("زيت الأرغان جُمان للشعر", "عناية بالشعر", "أرغان مغربي ١٠٠٪، للشعر الجاف والمتقصف، ٦٠ مل", "68 ريال"),
    ("نظارة شمسية أفق بإطار معدني", "إكسسوارات", "عدسات UV400 مستقطبة، إطار معدني خفيف، جراب صلب، للجنسين", "130 ريال"),
    ("طقم توابل نكهة ١٠ أنواع", "بقالة ومواد غذائية", "١٠ أنواع توابل مطحونة، عبوات زجاجية، للطبخ العربي", "110 ريال"),
    ("حذاء رياضي خطوة للجري", "أحذية رياضية", "نعل EVA مريح، شبك تهوية، وزن خفيف ٢٨٠ جرام، مقاسات ٣٩ إلى ٤٤", "220 ريال"),
    ("منظّم مكتب ترتيب خشبي", "أدوات مكتبية", "خشب MDF، ٥ أقسام، حامل أقلام وجوال، لون بني", "79 ريال"),
    ("شمعة معطّرة هدوء برائحة اللافندر", "ديكور منزلي", "شمع صويا طبيعي، رائحة لافندر، مدة احتراق ٤٠ ساعة، ٢٠٠ جرام", "60 ريال"),
    ("طقم عناية باللحية وقار", "عناية رجالية", "زيت لحية، بلسم، مشط خشبي، برائحة خشبية", "99 ريال"),
    ("دمية قطيفة أنيس دب ٤٠ سم", "ألعاب أطفال", "قطيفة ناعمة، حشو قطني، قابلة للغسل، مناسبة من الولادة", "70 ريال"),
    ("مظلة شمسية ظل محمولة", "مستلزمات خارجية", "حماية من الأشعة UV، فتح أوتوماتيكي، قطر ١٠٥ سم، حقيبة حمل", "55 ريال"),
    ("إبريق شاي زجاجي دفء ١ لتر", "المنزل والمطبخ", "زجاج بوروسيليكات، مصفاة ستانلس، سعة ١ لتر، مقاوم للحرارة", "88 ريال"),
    ("سوار لياقة نبض ذكي", "إلكترونيات", "قياس النبض والخطوات، شاشة AMOLED، بطارية ٧ أيام، مقاوم للماء IP68", "145 ريال"),
]

HEADERS = ["الاسم", "الفئة", "المميزات", "السعر"]

NAVY = "0E1118"
GOLD = "FFB800"


def main():
    out = sys.argv[1] if len(sys.argv) > 1 else "products.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "المنتجات"
    ws.sheet_view.rightToLeft = True

    thin = Side(style="thin", color="D0D0D0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    hfill = PatternFill("solid", fgColor=NAVY)
    hfont = Font(name="Arial", bold=True, size=12, color="FFFFFF")
    cfont = Font(name="Arial", size=11)

    for j, h in enumerate(HEADERS, 1):
        c = ws.cell(row=1, column=j, value=h)
        c.fill = hfill; c.font = hfont; c.border = border
        c.alignment = Alignment(horizontal="center", vertical="center")
    for i, row in enumerate(PRODUCTS, 2):
        for j, val in enumerate(row, 1):
            c = ws.cell(row=i, column=j, value=val)
            c.font = cfont; c.border = border
            c.alignment = Alignment(horizontal="right", vertical="center",
                                    wrap_text=(j == 3))
    widths = [34, 20, 52, 12]
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + j)].width = w
    ws.row_dimensions[1].height = 26
    for i in range(2, len(PRODUCTS) + 2):
        ws.row_dimensions[i].height = 34
    ws.freeze_panes = "A2"
    wb.save(out)
    print(f"[sample] wrote {out}  ({len(PRODUCTS)} products, {len(HEADERS)} columns)")


if __name__ == "__main__":
    main()
